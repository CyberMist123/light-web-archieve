"""docconv 的补充用例（现有 test_pdftext.py 已盖 PDF 文字层 / 扫描件 / 加密 / docx 基本形状）。

这里只碰不需要 OCR 的路径和 docx 细节：行内格式与超链接、表格转义与形状、坏文件、入口分流。
PDF 的 OCR 路径不测（现有测试已盖，且本文件不装 rapidocr 的假设要保持）。
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn

from link_brain import docconv, providers


def _para_with_runs(doc):
    p = doc.add_paragraph()
    a = p.add_run("加粗")
    a.bold = True
    b = p.add_run("继续加粗")
    b.bold = True                      # 相邻两个加粗 run
    p2 = doc.add_paragraph()
    it = p2.add_run("斜体词")
    it.italic = True
    p3 = doc.add_paragraph()
    rid = p3.part.rels.get_or_add_ext_rel(RT.HYPERLINK, "https://example.com/a")
    hl = p3._p.makeelement(qn("w:hyperlink"), {qn("r:id"): rid})
    r = hl.makeelement(qn("w:r"), {})
    t = r.makeelement(qn("w:t"), {})
    t.text = "仓库在这"
    r.append(t)
    hl.append(r)
    p3._p.append(hl)
    return p, p2, p3


def test_docx_inline_formatting_hyperlink_and_bold_merge(tmp_path):
    doc = Document()
    _para_with_runs(doc)
    path = tmp_path / "格式.docx"
    doc.save(str(path))
    out = docconv.docx_to_markdown(path)
    assert out["status"] == "ok" and out["method"] == "docx"
    md = out["markdown"]
    assert "**加粗继续加粗**" in md                     # 相邻加粗 run 的 **** 合并掉
    assert "*斜体词*" in md
    assert "[仓库在这](https://example.com/a)" in md


def test_docx_table_escaping_and_shapes(tmp_path):
    doc = Document()
    t1 = doc.add_table(rows=2, cols=2)
    t1.cell(0, 0).text = "名|字"
    t1.cell(0, 1).text = "说明\n换行"
    t1.cell(1, 0).text = "值"
    t1.cell(1, 1).text = "1"
    t2 = doc.add_table(rows=2, cols=1)                 # 单列表格 → 普通段落
    t2.cell(0, 0).text = "只有一列"
    t2.cell(1, 0).text = "第二行"
    t3 = doc.add_table(rows=2, cols=2)                 # 全空单元格表格（视觉方框）→ 跳过
    path = tmp_path / "表格.docx"
    doc.save(str(path))
    md = docconv.docx_markdown_text(path)
    assert "\\|" in md and ("说明 换行" in md)          # 竖线转义、单元格换行压成空格
    assert "只有一列\n\n第二行" in md                   # 单列表不当表格
    assert md.count("| --- | --- |") == 1              # 只有一张真表格的表头分隔行（空表跳过）


def test_docx_broken_bytes_is_unsupported(tmp_path):
    bad = tmp_path / "坏.docx"
    bad.write_bytes(b"not a zip at all")               # 存在但不是 docx（OXML 容器）
    out = docconv.docx_to_markdown(bad)
    assert out["status"] == "failed" and out["code"] == "PERMANENT.DOC_UNSUPPORTED"
    assert "打不开" in out["error"]


def test_open_pdf_missing_file_is_damaged(tmp_path):
    doc, err = docconv.open_pdf(tmp_path / "没有.pdf")
    assert doc is None and err["code"] == "PERMANENT.PDF_DAMAGED"


def test_to_markdown_unknown_suffix(monkeypatch, tmp_path):
    out = docconv.to_markdown(tmp_path / "图.png", cfg={"mode": "on"})
    assert out["status"] == "failed" and out["code"] == "PERMANENT.DOC_UNSUPPORTED"
    assert "png" in out["error"]


def test_to_markdown_skipped_when_off(monkeypatch, tmp_path):
    monkeypatch.setattr(providers, "resolve", lambda cap: None)
    monkeypatch.setattr(providers, "skipped_for", lambda cap: {"status": "skipped", "error": "功能关了"})
    out = docconv.to_markdown(tmp_path / "任何.docx", cfg=None)
    assert out["status"] == "skipped" and out["markdown"] is None
    assert out["method"] is None and out["note"] == "功能关了" and out["partial"] == 0
